"""
slide_builder.py
----------------
Assembles the PPTX deck from analytics outputs and chart images.
Uses python-pptx.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Literal

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

logger = logging.getLogger(__name__)

PerformanceSlideVariant = Literal["simple", "standard", "detailed"]
ExposureSlideVariant = Literal["net_bars", "pie_gross", "split", "table_only"]


def _hex_to_rgb(hex_str: str) -> RGBColor:
    return RGBColor(int(hex_str[:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


class DeckBuilder:
    """Builds an investor-ready PPTX slide deck."""

    def __init__(self, config: dict, metadata: dict | None = None):
        self.config = config
        self.brand = config["branding"]
        self._metadata = metadata or {}
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)

    def _brand_color(self, key: str, fallback: str) -> str:
        """Read brand colour from config with a safe fallback."""
        return self.brand.get(key, fallback)

    def _add_blank_slide(self, bg_color: str | None = None):
        layout = self.prs.slide_layouts[6]  # blank
        slide = self.prs.slides.add_slide(layout)
        if bg_color:
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = _hex_to_rgb(bg_color)
        return slide

    def _add_textbox(self, slide, left, top, width, height, text, font_size=14,
                      bold=False, color=None, font_name=None, alignment=PP_ALIGN.LEFT):
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = _hex_to_rgb(color or self.brand["text_dark"])
        p.font.name = font_name or self.brand["font_body"]
        p.alignment = alignment
        return txBox

    def _add_slide_footer(self, slide, *, bg: Literal["light", "dark"]) -> None:
        """As-of, generation time, and prices source (content slides)."""
        m = self._metadata
        if not m.get("as_of_iso"):
            return
        line1 = f"As of {m['as_of_iso']}  |  Generated {m.get('generated_at', '')}"
        line2 = m.get("prices_caption", "")
        text = f"{line1}\n{line2}" if line2 else line1
        color = self.brand["text_muted"] if bg == "light" else self.brand["secondary_color"]
        self._add_textbox(
            slide,
            Inches(0.9),
            Inches(6.95),
            Inches(11.5),
            Inches(0.52),
            text,
            font_size=8,
            color=color,
            alignment=PP_ALIGN.CENTER,
        )

    def add_title_slide(self, report_date: str, aum: str = "€125.4M"):
        slide = self._add_blank_slide(self.brand["bg_dark"])
        fund = self.config["fund"]

        # Gold accent line
        shape = slide.shapes.add_shape(
            1, Inches(1.5), Inches(2.8), Inches(2), Pt(3)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = _hex_to_rgb(self.brand["accent_color"])
        shape.line.fill.background()

        self._add_textbox(slide, Inches(1.5), Inches(3.0), Inches(10), Inches(1),
                          fund["name"], font_size=40, bold=True,
                          color=self._brand_color("text_light", "FFFFFF"), font_name=self.brand["font_heading"])
        self._add_textbox(slide, Inches(1.5), Inches(4.0), Inches(10), Inches(0.5),
                          fund["strategy"], font_size=18, color=self.brand["secondary_color"])
        self._add_textbox(slide, Inches(1.5), Inches(4.8), Inches(10), Inches(0.5),
                          f"Monthly Report  |  {report_date}  |  AUM: {aum}",
                          font_size=14, color=self.brand["secondary_color"])

        m = self._metadata
        if m.get("as_of_iso"):
            line1 = f"As of {m['as_of_iso']}  |  Generated {m.get('generated_at', '')}"
            line2 = m.get("prices_caption", "")
            block = f"{line1}\n{line2}" if line2 else line1
            self._add_textbox(
                slide,
                Inches(1.5),
                Inches(6.72),
                Inches(10.3),
                Inches(0.62),
                block,
                font_size=9,
                color=self.brand["secondary_color"],
                alignment=PP_ALIGN.CENTER,
            )

    def add_performance_slide(
        self,
        perf: dict,
        cum_chart_path: str,
        dd_chart_path: str,
        variant: PerformanceSlideVariant = "standard",
    ):
        slide = self._add_blank_slide(self.brand["bg_light"])
        self._add_textbox(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.6),
                          "Performance Summary", font_size=28, bold=True,
                          color=self.brand["primary_color"], font_name=self.brand["font_heading"])

        # KPI boxes
        x_start = 0.8
        for i, (label, value) in enumerate(perf.items()):
            x = Inches(x_start + i * 2.5)
            # Box background
            shape = slide.shapes.add_shape(1, x, Inches(1.3), Inches(2.2), Inches(1.0))
            shape.fill.solid()
            shape.fill.fore_color.rgb = _hex_to_rgb(self._brand_color("surface_color", "FFFFFF"))
            shape.line.color.rgb = _hex_to_rgb(self._brand_color("border_color", "E2E8F0"))
            shape.line.width = Pt(1)

            color = self._brand_color("positive_color", "22C55E") if value >= 0 else self._brand_color("negative_color", "EF4444")
            sign = "+" if value >= 0 else ""
            self._add_textbox(slide, x + Inches(0.15), Inches(1.35), Inches(2), Inches(0.6),
                              f"{sign}{value}%", font_size=26, bold=True, color=color)
            self._add_textbox(slide, x + Inches(0.15), Inches(1.85), Inches(2), Inches(0.3),
                              label, font_size=11, color=self.brand["text_muted"])

        if variant == "simple":
            self._add_slide_footer(slide, bg="light")
            return

        # Charts: standard = default sizes; detailed = taller charts for readability
        cum_top = Inches(2.6)
        cum_h = Inches(2.4 if variant == "detailed" else 2.2)
        dd_top = Inches(5.0 if variant == "standard" else 5.15)
        dd_h = Inches(2.2 if variant == "detailed" else 2.0)
        chart_w = Inches(11.5)

        if Path(cum_chart_path).exists():
            slide.shapes.add_picture(cum_chart_path, Inches(0.8), cum_top, chart_w, cum_h)
        if Path(dd_chart_path).exists():
            slide.shapes.add_picture(dd_chart_path, Inches(0.8), dd_top, chart_w, dd_h)

        self._add_slide_footer(slide, bg="light")

    def add_exposure_slide(
        self,
        title: str,
        chart_path: str,
        table_data: pd.DataFrame,
        variant: ExposureSlideVariant = "net_bars",
        secondary_chart_path: str | None = None,
    ):
        slide = self._add_blank_slide(self.brand["bg_light"])
        self._add_textbox(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.6),
                          title, font_size=28, bold=True,
                          color=self.brand["primary_color"], font_name=self.brand["font_heading"])

        if variant == "net_bars":
            if Path(chart_path).exists():
                slide.shapes.add_picture(chart_path, Inches(0.5), Inches(1.3), Inches(7), Inches(5.5))
        elif variant == "pie_gross":
            if Path(chart_path).exists():
                slide.shapes.add_picture(chart_path, Inches(2.0), Inches(1.2), Inches(5.5), Inches(5.6))
        elif variant == "split":
            if Path(chart_path).exists():
                slide.shapes.add_picture(chart_path, Inches(0.45), Inches(1.35), Inches(5.8), Inches(5.4))
            if secondary_chart_path and Path(secondary_chart_path).exists():
                slide.shapes.add_picture(secondary_chart_path, Inches(6.45), Inches(1.35), Inches(5.8), Inches(5.4))
            # No summary table: it sat in the same band as the pie and overlapped it on a standard slide.
            self._add_slide_footer(slide, bg="light")
            return
        # table_only: no chart — full-width summary table below title

        if variant == "table_only":
            table_left = Inches(0.8)
            table_width = Inches(11.5)
        else:
            # Summary table on the right (narrower when a chart occupies the left)
            table_left = Inches(8.0 if variant != "pie_gross" else 8.2)
            table_width = Inches(4.6)
        rows = min(len(table_data), 10)
        table_shape = slide.shapes.add_table(
            rows + 1, 3, table_left, Inches(1.3), table_width, Inches(0.4 * (rows + 1))
        )
        table = table_shape.table
        headers = ["", "Gross %", "Net %"]
        for j, h in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = h
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.bold = True
                p.font.color.rgb = _hex_to_rgb(self._brand_color("text_light", "FFFFFF"))
            cell.fill.solid()
            cell.fill.fore_color.rgb = _hex_to_rgb(self.brand["primary_color"])

        for i, (idx, row) in enumerate(table_data.head(rows).iterrows()):
            table.cell(i + 1, 0).text = str(idx)
            table.cell(i + 1, 1).text = f"{row['gross_weight']:.1f}%"
            table.cell(i + 1, 2).text = f"{row['net_weight']:+.1f}%"
            for j in range(3):
                for p in table.cell(i + 1, j).text_frame.paragraphs:
                    p.font.size = Pt(9)
                    p.font.color.rgb = _hex_to_rgb(self.brand["text_dark"])

        self._add_slide_footer(slide, bg="light")

    def add_top_positions_slide(self, longs: pd.DataFrame, shorts: pd.DataFrame):
        slide = self._add_blank_slide(self.brand["bg_light"])
        self._add_textbox(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.6),
                          "Top Positions", font_size=28, bold=True,
                          color=self.brand["primary_color"], font_name=self.brand["font_heading"])

        for col_idx, (label, df, color) in enumerate([
            ("Top Longs", longs, self._brand_color("positive_color", "22C55E")),
            ("Top Shorts", shorts, self._brand_color("negative_color", "EF4444")),
        ]):
            x_base = Inches(0.8 + col_idx * 6.2)
            self._add_textbox(slide, x_base, Inches(1.2), Inches(5), Inches(0.4),
                              label, font_size=16, bold=True, color=color)
            for i, (_, row) in enumerate(df.iterrows()):
                y = Inches(1.8 + i * 0.9)
                # Card
                shape = slide.shapes.add_shape(1, x_base, y, Inches(5.5), Inches(0.75))
                shape.fill.solid()
                shape.fill.fore_color.rgb = _hex_to_rgb(self._brand_color("surface_color", "FFFFFF"))
                shape.line.color.rgb = _hex_to_rgb(self._brand_color("border_color", "E2E8F0"))
                shape.line.width = Pt(1)

                self._add_textbox(slide, x_base + Inches(0.2), y + Inches(0.05), Inches(2), Inches(0.35),
                                  row["ticker"], font_size=14, bold=True)
                self._add_textbox(slide, x_base + Inches(0.2), y + Inches(0.38), Inches(2), Inches(0.3),
                                  f"{row['sector']}  •  {row['country']}", font_size=9, color=self.brand["text_muted"])

                pnl_color = self._brand_color("positive_color", "22C55E") if row["pnl"] >= 0 else self._brand_color("negative_color", "EF4444")
                pnl_sign = "+" if row["pnl"] >= 0 else ""
                self._add_textbox(slide, x_base + Inches(3.2), y + Inches(0.05), Inches(2), Inches(0.35),
                                  f"{row['weight_gross'] * 100:.1f}% weight", font_size=11,
                                  color=self.brand["text_muted"], alignment=PP_ALIGN.RIGHT)
                self._add_textbox(slide, x_base + Inches(3.2), y + Inches(0.38), Inches(2), Inches(0.3),
                                  f"{pnl_sign}{row['pnl_pct'] * 100:.1f}% P&L", font_size=10,
                                  color=pnl_color, alignment=PP_ALIGN.RIGHT)

        self._add_slide_footer(slide, bg="light")

    def add_risk_slide(self, metrics: dict):
        slide = self._add_blank_slide(self.brand["bg_light"])
        self._add_textbox(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.6),
                          "Risk Metrics", font_size=28, bold=True,
                          color=self.brand["primary_color"], font_name=self.brand["font_heading"])

        for i, (label, value) in enumerate(metrics.items()):
            col = i % 3
            row = i // 3
            x = Inches(0.8 + col * 4.0)
            y = Inches(1.5 + row * 2.5)

            shape = slide.shapes.add_shape(1, x, y, Inches(3.5), Inches(2.0))
            shape.fill.solid()
            shape.fill.fore_color.rgb = _hex_to_rgb(self._brand_color("surface_color", "FFFFFF"))
            shape.line.color.rgb = _hex_to_rgb(self._brand_color("border_color", "E2E8F0"))
            shape.line.width = Pt(1)

            display = f"{value}%" if "%" not in label and "Ratio" not in label else f"{value}"
            if "Ratio" in label:
                display = f"{value}x"
            self._add_textbox(slide, x + Inches(0.3), y + Inches(0.3), Inches(3), Inches(0.8),
                              str(display), font_size=36, bold=True, color=self.brand["primary_color"])
            self._add_textbox(slide, x + Inches(0.3), y + Inches(1.2), Inches(3), Inches(0.5),
                              label, font_size=12, color=self.brand["text_muted"])

        self._add_slide_footer(slide, bg="light")

    def add_attribution_slide(self, attribution: pd.DataFrame, chart_path: str):
        slide = self._add_blank_slide(self.brand["bg_light"])
        self._add_textbox(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.6),
                          "P&L Attribution", font_size=28, bold=True,
                          color=self.brand["primary_color"], font_name=self.brand["font_heading"])

        if Path(chart_path).exists():
            slide.shapes.add_picture(chart_path, Inches(0.6), Inches(1.3), Inches(7.0), Inches(5.6))

        rows = min(len(attribution), 10)
        if rows == 0:
            self._add_textbox(
                slide, Inches(8.0), Inches(2.8), Inches(4.8), Inches(1.0),
                "No attribution data available", font_size=14, color=self.brand["text_muted"]
            )
            self._add_slide_footer(slide, bg="light")
            return

        table_shape = slide.shapes.add_table(
            rows + 1, 3, Inches(8.0), Inches(1.3), Inches(4.8), Inches(0.42 * (rows + 1))
        )
        table = table_shape.table
        headers = ["Group", "P&L", "Contribution %"]

        for j, header in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = _hex_to_rgb(self.brand["primary_color"])
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.bold = True
                p.font.color.rgb = _hex_to_rgb(self._brand_color("text_light", "FFFFFF"))

        for i, (group_name, row) in enumerate(attribution.head(rows).iterrows()):
            table.cell(i + 1, 0).text = str(group_name)
            table.cell(i + 1, 1).text = f"{row['total_pnl']:,.0f}"
            contrib = row["pnl_contribution_pct"]
            table.cell(i + 1, 2).text = f"{contrib:+.2f}%"

            for j in range(3):
                cell = table.cell(i + 1, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _hex_to_rgb(self._brand_color("surface_color", "FFFFFF"))
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(9)
                    p.font.color.rgb = _hex_to_rgb(self.brand["text_dark"])

            contrib_color = self._brand_color("positive_color", "22C55E") if contrib >= 0 else self._brand_color("negative_color", "EF4444")
            for p in table.cell(i + 1, 2).text_frame.paragraphs:
                p.font.color.rgb = _hex_to_rgb(contrib_color)
                p.alignment = PP_ALIGN.RIGHT

        self._add_slide_footer(slide, bg="light")

    def add_disclaimer_slide(self):
        slide = self._add_blank_slide(self.brand["bg_dark"])
        self._add_textbox(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(0.6),
                          self.config["fund"]["legal_name"], font_size=24, bold=True,
                          color=self._brand_color("text_light", "FFFFFF"), font_name=self.brand["font_heading"])
        self._add_textbox(slide, Inches(1.5), Inches(2.5), Inches(10), Inches(4),
                          self.config["disclaimer_text"], font_size=11, color=self.brand["secondary_color"])
        self._add_slide_footer(slide, bg="dark")

    def save(self, output_path: str):
        self.prs.save(output_path)
        logger.info("Deck saved to %s", output_path)
